from pathlib import Path
from typing import Any

from docx import Document
from pptx import Presentation
from openpyxl import Workbook


class FileExportService:
    """
    Export Nova AI reports into DOCX, PPTX and XLSX files.
    """

    def __init__(
        self,
        output_dir: str = "generated/exports",
    ):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

    def _safe_name(self, report: dict[str, Any]) -> str:
        query = str(
            report.get(
                "query",
                "nova_ai_report",
            )
        ).strip()

        safe_name = "".join(
            character
            if character.isalnum()
            else "_"
            for character in query
        ).strip("_")

        return safe_name or "nova_ai_report"

    def _versioned_file_path(
        self,
        report: dict[str, Any],
        extension: str,
    ) -> Path:
        """
        Create a unique versioned file path.

        Example:
        AI_Workspace_v1.docx
        AI_Workspace_v2.docx
        AI_Workspace_v3.docx
        """

        safe_name = self._safe_name(report)

        version = 1

        while True:
            file_path = (
                self.output_dir
                / f"{safe_name}_v{version}{extension}"
            )

            if not file_path.exists():
                return file_path

            version += 1

    def _get_title(self, report: dict[str, Any]) -> str:
        return str(
            report.get(
                "title",
                "Nova AI Report",
            )
        )

    def _get_summary(self, report: dict[str, Any]) -> str:
        return str(
            report.get(
                "summary",
                "",
            )
        )

    def _get_findings(
        self,
        report: dict[str, Any],
    ) -> list[Any]:
        findings = report.get(
            "key_findings",
            [],
        )

        return findings if isinstance(
            findings,
            list,
        ) else []

    def _get_citations(
        self,
        report: dict[str, Any],
    ) -> list[Any]:
        citations = report.get(
            "citations",
            [],
        )

        return citations if isinstance(
            citations,
            list,
        ) else []

    # =====================================================
    # DOCX
    # =====================================================

    def export_docx(
        self,
        report: dict[str, Any],
    ) -> str:

        if not report:
            raise ValueError(
                "Report cannot be empty."
            )

        file_path = self._versioned_file_path(
            report,
            ".docx",
        )

        document = Document()

        document.add_heading(
            self._get_title(report),
            level=0,
        )

        document.add_heading(
            "Query",
            level=1,
        )

        document.add_paragraph(
            str(
                report.get(
                    "query",
                    "",
                )
            )
        )

        document.add_heading(
            "Summary",
            level=1,
        )

        document.add_paragraph(
            self._get_summary(report)
        )

        document.add_heading(
            "Key Findings",
            level=1,
        )

        findings = self._get_findings(report)

        if findings:
            for finding in findings:
                document.add_paragraph(
                    str(finding),
                    style="List Bullet",
                )
        else:
            document.add_paragraph(
                "No key findings available."
            )

        document.add_heading(
            "Sources & Citations",
            level=1,
        )

        citations = self._get_citations(report)

        if citations:
            for citation in citations:

                if not isinstance(
                    citation,
                    dict,
                ):
                    document.add_paragraph(
                        str(citation)
                    )
                    continue

                citation_id = citation.get(
                    "id",
                    "",
                )

                title = citation.get(
                    "title",
                    "Web Source",
                )

                source_type = citation.get(
                    "source_type",
                    "web",
                )

                url = citation.get(
                    "url",
                    "",
                )

                document.add_paragraph(
                    f"[{citation_id}] {title}"
                )

                document.add_paragraph(
                    f"Type: {source_type}"
                )

                if url:
                    document.add_paragraph(
                        f"URL: {url}"
                    )
        else:
            document.add_paragraph(
                "No sources available."
            )

        document.save(
            str(file_path)
        )

        return str(file_path)

    # =====================================================
    # PPTX
    # =====================================================

    def export_pptx(
        self,
        report: dict[str, Any],
    ) -> str:

        if not report:
            raise ValueError(
                "Report cannot be empty."
            )

        file_path = self._versioned_file_path(
            report,
            ".pptx",
        )

        presentation = Presentation()

        # Title slide
        title_slide = presentation.slides.add_slide(
            presentation.slide_layouts[0]
        )

        title_slide.shapes.title.text = (
            self._get_title(report)
        )

        title_slide.placeholders[1].text = (
            str(
                report.get(
                    "query",
                    "",
                )
            )
        )

        # Summary slide
        summary_slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        summary_slide.shapes.title.text = "Summary"

        summary_slide.placeholders[1].text = (
            self._get_summary(report)
            or "No summary available."
        )

        # Findings slide
        findings_slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        findings_slide.shapes.title.text = (
            "Key Findings"
        )

        text_frame = (
            findings_slide.placeholders[1].text_frame
        )

        text_frame.clear()

        findings = self._get_findings(report)

        if findings:
            for index, finding in enumerate(
                findings
            ):
                paragraph = (
                    text_frame.paragraphs[0]
                    if index == 0
                    else text_frame.add_paragraph()
                )

                paragraph.text = (
                    f"{index + 1}. {finding}"
                )
        else:
            text_frame.text = (
                "No key findings available."
            )

        # Sources slide
        sources_slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        sources_slide.shapes.title.text = (
            "Sources & Citations"
        )

        source_frame = (
            sources_slide.placeholders[1].text_frame
        )

        source_frame.clear()

        citations = self._get_citations(report)

        if citations:

            for index, citation in enumerate(
                citations
            ):

                if not isinstance(
                    citation,
                    dict,
                ):
                    text = str(citation)
                else:
                    citation_id = citation.get(
                        "id",
                        "",
                    )

                    title = citation.get(
                        "title",
                        "Web Source",
                    )

                    url = citation.get(
                        "url",
                        "",
                    )

                    text = (
                        f"[{citation_id}] {title}"
                    )

                    if url:
                        text += f"\n{url}"

                paragraph = (
                    source_frame.paragraphs[0]
                    if index == 0
                    else source_frame.add_paragraph()
                )

                paragraph.text = text

        else:
            source_frame.text = (
                "No sources available."
            )

        presentation.save(
            str(file_path)
        )

        return str(file_path)

    # =====================================================
    # XLSX
    # =====================================================

    def export_xlsx(
        self,
        report: dict[str, Any],
    ) -> str:

        if not report:
            raise ValueError(
                "Report cannot be empty."
            )

        file_path = self._versioned_file_path(
            report,
            ".xlsx",
        )

        workbook = Workbook()

        worksheet = workbook.active
        worksheet.title = "Research Report"

        worksheet["A1"] = "Title"
        worksheet["B1"] = self._get_title(report)

        worksheet["A2"] = "Query"
        worksheet["B2"] = str(
            report.get(
                "query",
                "",
            )
        )

        worksheet["A4"] = "Summary"
        worksheet["B4"] = self._get_summary(report)

        # Key findings
        worksheet["A6"] = "Key Findings"

        findings = self._get_findings(report)

        for index, finding in enumerate(
            findings,
            start=7,
        ):
            worksheet[f"A{index}"] = index - 6
            worksheet[f"B{index}"] = str(finding)

        # Citations
        citation_start = max(
            7 + len(findings),
            9,
        )

        worksheet[
            f"A{citation_start}"
        ] = "Sources & Citations"

        worksheet[
            f"A{citation_start + 1}"
        ] = "ID"

        worksheet[
            f"B{citation_start + 1}"
        ] = "Title"

        worksheet[
            f"C{citation_start + 1}"
        ] = "Type"

        worksheet[
            f"D{citation_start + 1}"
        ] = "URL"

        citations = self._get_citations(report)

        for index, citation in enumerate(
            citations,
            start=citation_start + 2,
        ):

            if not isinstance(
                citation,
                dict,
            ):
                worksheet[
                    f"B{index}"
                ] = str(citation)
                continue

            worksheet[
                f"A{index}"
            ] = citation.get(
                "id",
                "",
            )

            worksheet[
                f"B{index}"
            ] = citation.get(
                "title",
                "Web Source",
            )

            worksheet[
                f"C{index}"
            ] = citation.get(
                "source_type",
                "web",
            )

            worksheet[
                f"D{index}"
            ] = citation.get(
                "url",
                "",
            )

        # Basic column sizing
        worksheet.column_dimensions[
            "A"
        ].width = 20

        worksheet.column_dimensions[
            "B"
        ].width = 70

        worksheet.column_dimensions[
            "C"
        ].width = 20

        worksheet.column_dimensions[
            "D"
        ].width = 80

        workbook.save(
            str(file_path)
        )

        return str(file_path)


def get_file_export_service():
    return FileExportService()